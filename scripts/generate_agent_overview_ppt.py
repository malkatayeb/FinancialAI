"""One-off generator: Financial AI agent overview deck. Run: python scripts/generate_agent_overview_ppt.py"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "FinancialAI-Agent-Overview.pptx"


def add_title_slide(prs, title: str, subtitle: str):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets(prs, title: str, lines: list[str]):
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = lines[0]
    tf.paragraphs[0].font.size = Pt(20)
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(20)


def add_sample_response_slide(prs):
    """Title-only layout + textbox so full JSON / reply fits."""
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Sample response (HTTP 200)"
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)

    reply = (
        "Your current account balance is $12,602.51.\\n\\n"
        "Here are your recent Grocery transactions:\\n\\n"
        "* Transaction 1: $25.00 (Purchased milk and bread)\\n"
        "* Transaction 2: $15.00 (Bought eggs and cheese)\\n\\n"
        "Total Groceries transaction amount: $40.00\\n\\n"
        "According to your monthly budget, you have allocated $50.00 for groceries. "
        "Unfortunately, this month's transactions ($40.00) are short of the budget by $10.00. "
        "You are currently under the groceries monthly budget by $10.00."
    )
    body = (
        "Same shape as ChatResponse: sessionId + reply. "
        "Example JSON (\\n shown explicitly as in a raw JSON string):\n\n"
        "{\n"
        f'  "reply": "{reply}",\n'
        '  "sessionId": "default"\n'
        "}"
    )

    left, top, w, h = Inches(0.45), Inches(1.05), Inches(9.1), Inches(5.95)
    box = slide.shapes.add_textbox(left, top, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    p.font.size = Pt(11)
    p.font.name = "Consolas"


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Financial AI Agent",
        "Bank Discount — capabilities, build, and runtime behavior",
    )

    add_bullets(
        prs,
        "Capabilities",
        [
            "Answers in natural language about money: balance, spending by category, budget vs limit.",
            "Uses real data: tools run JPA queries on the in-memory H2 database (seeded transactions & budgets).",
            "Multi-turn memory: optional sessionId isolates conversation context (default session if omitted).",
            "Simple API: POST or GET /api/chat — one endpoint for all intents (no hand-coded message routing).",
        ],
    )

    add_bullets(
        prs,
        "How we built it",
        [
            "Java 21 + Spring Boot 3.3 (REST, JPA, H2).",
            "LangChain4j: @AiService advisor + @Tool methods on a Spring @Service bean.",
            "Ollama: local chat model (e.g. llama3.2) via langchain4j-ollama-spring-boot-starter + application.yml.",
            "Docker Compose option: app + Ollama + model pull; or run Maven locally against localhost:11434.",
        ],
    )

    add_bullets(
        prs,
        "What it is doing (request flow)",
        [
            "1. Single entry — every chat calls financialAgent.chat(memoryId, message) from /api/chat.",
            "2. Tools wired — @AiService(tools = \"bankingService\") exposes @Tool names, descriptions, parameters to the LLM.",
            "3. Model decides — Ollama receives system + user text + tool definitions; it replies with tool calls or final text.",
            "4. Java executes — LangChain4j invokes BankingService methods (not arbitrary HTTP); results return to the model.",
            "5. Loop until done — the model may call more tools, then returns the final reply string to the client.",
        ],
    )

    add_bullets(
        prs,
        "Registered tools (examples)",
        [
            "getBalance — sum of posted transactions (credits minus debits).",
            "getTransactionsByCategory(category) — recent rows for that category.",
            "checkBudgetStatus(category) — spend vs monthly budget for the reporting month; exact over/under.",
            "The LLM chooses tools from these descriptions; your code does not keyword-match user text.",
        ],
    )

    add_bullets(
        prs,
        "Run & try it",
        [
            "Docker: docker compose up --build (app on :8080, Ollama on :11434).",
            "Local: ollama pull llama3.2, then mvn spring-boot:run from project root.",
            "POST http://localhost:8080/api/chat — sample JSON (matches seeded H2 data):",
            '{"message":"What is my current account balance? Show my recent Groceries transactions and say if I am under or over the Groceries monthly budget.","sessionId":"demo-1"}',
        ],
    )

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
